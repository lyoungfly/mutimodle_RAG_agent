"""从已校验结果渲染报告；不请求模型，不执行文档中的链接或标记。"""

import io
import json
import hashlib


FORMATS = {
    "json": "application/json",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}


def _text(value) -> str:
    text = value if isinstance(value, str) else str(value)
    return "".join(char if char in "\n\t" or (
        ord(char) >= 32 and not 0xD800 <= ord(char) <= 0xDFFF
        and ord(char) not in {0xFFFE, 0xFFFF}
    ) else "\ufffd" for char in text)


def _dump(value) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2, allow_nan=False)


def _references(ids) -> str:
    return " ".join(f"[{item}]" for item in ids)


def _row_ranges(rows) -> str:
    """连续原始行号压缩展示；保留不连续行，不能误称为一个连续范围。"""
    if not rows:
        return "无选中行"
    values = sorted(set(rows))
    parts, first, last = [], values[0], values[0]
    for value in values[1:]:
        if value == last + 1:
            last = value
        else:
            parts.append(str(first) if first == last else f"{first}–{last}")
            first = last = value
    parts.append(str(first) if first == last else f"{first}–{last}")
    return ", ".join(parts)


def _analysis_report(calculation: dict) -> dict:
    """独立表格任务同样可导出；所有文字直接来自计算结果，不经过模型。"""
    calculation_id = calculation["calculation_id"]
    evidence_id = "e_" + hashlib.sha256(calculation_id.encode()).hexdigest()[:24]
    parameters = calculation.get("parameters", {})
    summary = calculation["summary"]
    operation = summary["operation"]
    labels = {"count": "计数", "sum": "求和", "mean": "平均值", "min": "最小值",
              "max": "最大值", "group_count": "分组计数", "group_sum": "分组求和"}
    value = summary.get("value")
    result_text = "无可计算数值" if value is None else str(value)
    text = (f"{calculation['filename']} / {calculation['table_id']}："
            f"{labels.get(operation, operation)}结果为 {result_text}，"
            f"选中 {summary['selected_rows']} 行，共 {summary['total_rows']} 行。")
    document_id = calculation["document_id"]
    return {
        "question": f"对 {calculation['filename']} 的 {calculation['table_id']} 执行受控表格分析。",
        "collection_id": calculation.get("collection_id", ""),
        "completed_at": calculation.get("completed_at", "见任务记录"),
        "summary": text, "summary_evidence_ids": [evidence_id],
        "findings": [{"text": text, "evidence_ids": [evidence_id]}],
        "conflicts": [], "recommendations": [], "missing_information": [],
        "document_ids": [document_id], "calculations": [calculation],
        "evidence": [{
            "evidence_id": evidence_id, "kind": "calculation", "calculation_id": calculation_id,
            "document_id": document_id, "filename": calculation["filename"],
            "source": calculation.get("source", {}), "content": _dump(summary),
        }],
        "validation_note": "数值由受控表格工具计算；报告保留参数、原始行号、数据快照摘要及处理规则。",
    }


def _blocks(result):
    yield 0, "IndustrialInsight 企业分析报告"
    yield 3, f"生成时间（UTC）：{result.get('completed_at', '')}"
    yield 3, f"知识库：{result.get('collection_id', '')}"
    yield 1, "任务与分析范围"
    yield 3, result["question"]
    yield 3, f"涉及 {len(result['document_ids'])} 份资料；报告保留本次任务的证据和计算快照。"
    yield 1, "摘要"
    yield 3, result["summary"] + " " + _references(result["summary_evidence_ids"])
    yield 1, "主要发现"
    for number, item in enumerate(result["findings"], 1):
        yield 3, f"{number}. {item['text']} {_references(item['evidence_ids'])}"
    if not result["findings"]:
        yield 3, "暂无具备足够证据的事实结论。"
    for key, heading in [("conflicts", "资料差异与冲突"), ("recommendations", "分析建议（待业务确认）"),
                         ("missing_information", "资料缺口")]:
        yield 1, heading
        if not result[key]:
            yield 3, "本次分析未列出。"
        for number, item in enumerate(result[key], 1):
            text = item if isinstance(item, str) else item["text"] + " " + _references(item["evidence_ids"])
            yield 3, f"{number}. {text}"
    yield 1, "计算记录"
    if not result["calculations"]:
        yield 3, "本次任务未执行表格计算。"
    for number, item in enumerate(result["calculations"], 1):
        yield 2, f"计算 {number}：{item['filename']} / {item['table_id']}"
        yield 3, f"计算编号：{item['calculation_id']}"
        yield 3, "参数与筛选条件：\n" + _dump(item.get("parameters", {}))
        yield 3, "计算结果：\n" + _dump(item.get("summary", {}))
        for group in item.get("groups", []):
            yield 3, f"分组 {group.get('label', group.get('key', ''))}：值 {group.get('value')}；行数 {group.get('count')}"
            yield 3, "该组源行：" + _row_ranges(group.get("row_numbers", []))
        yield 3, "全部选中源行：" + _row_ranges(item.get("selected_row_numbers", []))
        yield 3, "来源与快照：\n" + _dump(item.get("snapshot", {}))
        yield 3, "数值与缺失值处理规则：\n" + _dump(item.get("policy", {}))
        for warning in item.get("warnings", []):
            yield 3, "提示：" + str(warning)
        yield 3, "数据预览（仅为有限预览，统计使用全部选中行）：\n" + _dump(item.get("rows", []))
    yield 1, "引用证据快照"
    for item in result["evidence"]:
        yield 2, f"[{item['evidence_id']}] {item['filename']}"
        yield 3, f"文档 ID：{item['document_id']}"
        if item.get("document_version"):
            yield 3, "解析文档快照 SHA-256：" + item["document_version"]
        location = {key: item[key] for key in ("page", "section", "chunk_id", "calculation_id") if item.get(key)}
        yield 3, "定位：" + _dump({**location, "source": item.get("source", {})})
        url = item.get("url") or item.get("source", {}).get("file_url", "")
        if url:
            yield 3, "应用内原文地址：" + url
        yield 3, item.get("content", "")
        if item.get("content_truncated"):
            yield 3, "本条证据展示已限长，完整内容请打开原文。"
    yield 1, "报告说明"
    yield 3, result.get("validation_note", "已校验引用编号；事实与建议应结合原文复核。")
    yield 3, "本报告由同一份结构化结果直接渲染。引用地址需要登录原应用并具备资料权限。资料变更后，报告仍展示生成时的快照。"


def _render_docx(blocks) -> bytes:
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError('DOCX 导出依赖未安装，请执行 python -m pip install -e ".[office]"。') from exc
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    for name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.color.rgb = RGBColor.from_string("155B4B")
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    for level, text in blocks:
        if level < 3:
            document.add_heading(text, level=level)
        else:
            document.add_paragraph(text)
    footer = section.footer.paragraphs[0]
    footer.add_run("IndustrialInsight  |  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _render_pdf(blocks) -> bytes:
    try:
        import pymupdf
    except ImportError as exc:
        raise RuntimeError('PDF 导出依赖未安装，请执行 python -m pip install -e ".[pdf]"。') from exc
    # 使用 MuPDF 自带的中文字体；不读取外部 HTML、字体网址或文档链接。
    font = pymupdf.Font("cjk")
    widths = {}
    with pymupdf.open() as document:
        page = None
        y = 0

        def new_page():
            if len(document) >= 200:
                raise ValueError("PDF 超过 200 页，请拆分分析任务或下载 JSON 快照。")
            current = document.new_page(width=595, height=842)
            current.insert_font(fontname="report-cjk", fontbuffer=font.buffer)
            current.insert_text((46, 28), "IndustrialInsight 企业分析报告", fontname="report-cjk",
                                fontsize=8, color=(0.3, 0.4, 0.37))
            current.insert_text((270, 819), f"{len(document)}", fontsize=8)
            return current, 58

        for level, text in blocks:
            size = [19, 14, 11.5, 10][level]
            line_height = size * 1.65
            color = (0.08, 0.33, 0.27) if level < 3 else (0.15, 0.18, 0.17)
            if page is None or y + line_height * 2 > 792:
                page, y = new_page()
            elif level < 3:
                y += 8
            for raw_line in text.expandtabs(4).split("\n"):
                line, width = "", 0.0
                lines = []
                for character in raw_line:
                    key = (character, size)
                    if key not in widths:
                        widths[key] = font.text_length(character, fontsize=size)
                    advance = widths[key]
                    if line and width + advance > 503:
                        lines.append(line)
                        line, width = "", 0.0
                    line += character
                    width += advance
                lines.append(line)
                for line in lines:
                    if y + line_height > 792:
                        page, y = new_page()
                    if line:
                        page.insert_text((46, y), line, fontname="report-cjk",
                                         fontsize=size, color=color)
                    y += line_height
            y += 4
        document.set_metadata({"title": "IndustrialInsight 企业分析报告", "author": "IndustrialInsight"})
        return document.tobytes(garbage=4, deflate=True)


def render_report(result: dict, format: str = "docx") -> bytes:
    from papermind.enterprise.agent_schema import validate_report

    if format not in FORMATS:
        raise ValueError("报告格式仅支持 docx、pdf、json。")
    if isinstance(result, dict) and "calculation_id" in result and "findings" not in result:
        result = _analysis_report(result)
    validated = validate_report(result)
    if format == "json":
        return _dump(validated).encode("utf-8")
    blocks = [(level, _text(text)) for level, text in _blocks(validated)]
    if sum(len(text) for _, text in blocks) > 600000:
        raise ValueError("报告正文超过 60 万字符，请拆分分析任务或下载 JSON 快照。")
    return _render_docx(blocks) if format == "docx" else _render_pdf(blocks)
