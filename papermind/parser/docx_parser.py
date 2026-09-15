import re
from pathlib import Path

from papermind.models import Document, Paragraph, Section, Table


def parse_docx(path: Path, document: Document) -> Document:
    from docx import Document as WordDocument
    from docx.text.paragraph import Paragraph as WordParagraph

    word = WordDocument(path)
    document.metadata.update(
        format="docx",
        warnings=[
            "Word 按正文段落和表格定位，不提供分页；页眉页脚、修订内容、文本框和内嵌图片暂不解析。"
        ],
    )
    current = Section(document.title, 1)
    document.sections.append(current)
    parents = []
    paragraph_number = 0
    cell_count = 0
    for block in word.iter_inner_content():
        if isinstance(block, WordParagraph):
            paragraph_number += 1
            if paragraph_number > 20000:
                raise ValueError("Word 正文超过 20000 段，请拆分后上传。")
            text = block.text.strip()
            if not text:
                continue
            heading = (
                re.fullmatch(
                    r"(?:Heading|标题)\s*(\d)", block.style.name or "", re.IGNORECASE
                )
                if block.style
                else None
            )
            if heading:
                level = int(heading[1])
                while parents and parents[-1][0] >= level:
                    parents.pop()
                parents.append((level, text))
                current = Section(text, level)
                document.sections.append(current)
            current.paragraphs.append(
                Paragraph(
                    text,
                    0,
                    source={
                        "format": "docx",
                        "paragraph": paragraph_number,
                        "heading_path": " / ".join(title for _, title in parents),
                    },
                )
            )
        else:
            rows = []
            for row in block.rows:
                cell_count += len(row.cells)
                if cell_count > 100000:
                    raise ValueError("Word 表格超过 100000 个单元格，请拆分后上传。")
                rows.append([cell.text.strip() for cell in row.cells])
            width = max(map(len, rows), default=0)
            if not width:
                continue
            number = len(document.tables) + 1
            document.tables.append(
                Table(
                    f"table_{number}",
                    0,
                    [f"列 {i + 1}" for i in range(width)],
                    [row + [""] * (width - len(row)) for row in rows],
                    caption=f"表格 {number}",
                    section=current.title,
                    source={
                        "format": "docx",
                        "table_number": number,
                        "heading_path": " / ".join(title for _, title in parents),
                    },
                    row_numbers=list(range(1, len(rows) + 1)),
                )
            )
    return document
