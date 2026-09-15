from contextlib import ExitStack
from datetime import date, datetime, time
from itertools import zip_longest
from pathlib import Path

from papermind.models import Document, Table


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def parse_xlsx(path: Path, document: Document) -> Document:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    warnings = [
        "每张工作表首个非空行作为表头；隐藏工作表跳过，图表和内嵌图片暂不解析。"
    ]
    document.metadata.update(format="xlsx", warnings=warnings)
    with ExitStack() as stack:
        values = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        stack.callback(values.close)
        formulas = load_workbook(
            path, read_only=True, data_only=False, keep_links=False
        )
        stack.callback(formulas.close)
        if len(values.sheetnames) > 100:
            raise ValueError("Excel 超过 100 张工作表，请拆分后上传。")
        count = missing = formula_count = 0
        for sheet in values.worksheets:
            if sheet.sheet_state != "visible":
                continue
            formula_sheet = formulas[sheet.title]
            # 不信任文件声明的 dimensions；坐标上限已在 ZIP/XML 校验时检查。
            sheet.reset_dimensions()
            formula_sheet.reset_dimensions()
            rows, row_numbers = [], []
            header_row = 0
            for row_number, pair in enumerate(
                zip_longest(sheet.iter_rows(), formula_sheet.iter_rows()), 1
            ):
                cached, original = pair
                if cached is None or original is None or row_number > 20000:
                    raise ValueError("Excel 工作表结构不一致或超过行数限制。")
                count += max(len(cached), len(original))
                if count > 100000:
                    raise ValueError("Excel 超过 100000 个单元格，请拆分后上传。")
                row = []
                for cell, formula in zip_longest(cached, original):
                    value = cell.value if cell else None
                    if formula is not None and formula.data_type == "f":
                        formula_count += 1
                        if value is None:
                            missing += 1
                            value = "[公式结果不可用]"
                    row.append(cell_text(value))
                while row and not row[-1]:
                    row.pop()
                if not row:
                    continue
                if not header_row:
                    header_row, headers = row_number, row
                else:
                    rows.append(row)
                    row_numbers.append(row_number)
            if not header_row:
                continue
            width = max([len(headers), *(len(row) for row in rows)])
            headers += [f"列 {i + 1}" for i in range(len(headers), width)]
            headers = [h or f"列 {i + 1}" for i, h in enumerate(headers)]
            end = get_column_letter(width)
            document.tables.append(
                Table(
                    f"table_{len(document.tables) + 1}",
                    0,
                    headers,
                    [row + [""] * (width - len(row)) for row in rows],
                    caption=sheet.title,
                    section=sheet.title,
                    source={
                        "format": "xlsx",
                        "sheet": sheet.title,
                        "header_row": header_row,
                        "last_column": end,
                        "cell_range": f"A{header_row}:{end}{row_numbers[-1] if rows else header_row}",
                    },
                    row_numbers=row_numbers,
                )
            )
        if formula_count:
            warnings.append(
                "公式仅使用文件保存时的缓存结果，不执行计算；请在 Excel 中重算并保存后上传。"
            )
        if missing:
            warnings.append(
                f"{missing} 个公式缺少缓存结果，已标记为不可用，未作为数值处理。"
            )
    return document
